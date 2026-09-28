"""Recalculate every built workbook with LibreOffice and check the results.

    python verify.py            # needs `soffice` (LibreOffice Calc) on PATH

1. Scans every formula cell for spreadsheet errors (#NAME?, #VALUE!, #REF!, ...).
2. Re-computes the key numbers in plain Python from the demo data and
   compares them with what the spreadsheet calculated.
"""

import collections
import glob
import os
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

import build_workbook as B

HERE = os.path.dirname(os.path.abspath(__file__))
DIST = os.path.join(HERE, "..", "dist")
ERRORS = ("#NAME?", "#VALUE!", "#REF!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:")
failures = []


def check(cond, msg):
    if not cond:
        failures.append(msg)
        print("  FAIL:", msg)


def close(a, b, tol=0.011):
    return a is not None and b is not None and abs(float(a) - float(b)) <= tol


def recalc(paths, outdir):
    """Open each file in LibreOffice, force a full hard recalculation, save a copy.

    A plain `soffice --convert-to` does not always recalculate long dependency
    chains in order for files that have no cached values, so we drive
    LibreOffice through UNO and call calculateAll() explicitly.
    """
    import time
    import uno
    from com.sun.star.beans import PropertyValue

    def prop(name, value):
        p = PropertyValue()
        p.Name, p.Value = name, value
        return p

    port = 2002
    proc = subprocess.Popen(["soffice", "--headless", "--invisible", "--norestore",
                             f"--accept=socket,host=127.0.0.1,port={port};urp;"],
                            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        local = uno.getComponentContext()
        resolver = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
        for _ in range(60):
            try:
                ctx = resolver.resolve(f"uno:socket,host=127.0.0.1,port={port};urp;StarOffice.ComponentContext")
                break
            except Exception:
                time.sleep(0.5)
        else:
            raise RuntimeError("could not connect to LibreOffice")
        desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
        for path in paths:
            url = uno.systemPathToFileUrl(os.path.abspath(path))
            doc = desktop.loadComponentFromURL(url, "_blank", 0, (prop("Hidden", True),))
            doc.calculateAll()
            doc.calculateAll()
            out = uno.systemPathToFileUrl(os.path.join(os.path.abspath(outdir), os.path.basename(path)))
            doc.storeToURL(out, (prop("FilterName", "Calc Office Open XML"),))
            doc.close(True)
        try:
            desktop.terminate()
        except Exception:
            pass
    finally:
        try:
            proc.wait(timeout=20)
        except subprocess.TimeoutExpired:
            proc.kill()


def scan_errors(src_path, calc_path):
    f = load_workbook(src_path)  # formulas
    v = load_workbook(calc_path, data_only=True)  # values
    n_formulas = 0
    for ws in f.worksheets:
        wv = v[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n_formulas += 1
                    val = wv[c.coordinate].value
                    if isinstance(val, str) and val.startswith(ERRORS):
                        check(False, f"{os.path.basename(src_path)} {ws.title}!{c.coordinate} = {val}  ({c.value[:80]})")
    return n_formulas


def simulate_debts(debts, extra, strategy):
    """Plain-Python snowball/avalanche, same rules as the spreadsheet."""
    if strategy == "Avalanche":
        order = sorted(range(len(debts)), key=lambda i: (-debts[i][2], i))
    else:
        order = sorted(range(len(debts)), key=lambda i: (debts[i][1], i))
    bal = [debts[i][1] for i in order]
    apr = [debts[i][2] for i in order]
    mins = [debts[i][3] for i in order]
    budget = sum(mins) + extra
    paid = [0.0] * len(bal)
    payoff = [None] * len(bal)
    for m in range(1, B.SCHED_MONTHS + 1):
        owed = [0 if b <= 0.004 else round(b * (1 + a / 1200), 2) for b, a in zip(bal, apr)]
        base = [min(mn, o) for mn, o in zip(mins, owed)]
        left = budget - sum(base)
        ext = []
        for o, b0 in zip(owed, base):
            e = max(0, min(o - b0, left - sum(ext)))
            ext.append(e)
        for k in range(len(bal)):
            pay = base[k] + ext[k]
            paid[k] += pay
            nb = max(0, round(owed[k] - pay, 2))
            if nb <= 0.004 and bal[k] > 0.004 and payoff[k] is None:
                payoff[k] = m
            bal[k] = nb
    names = [debts[i][0] for i in order]
    return {names[k]: (payoff[k], round(paid[k], 2)) for k in range(len(names))}


def verify_bundle_demo(calc_path):
    wb = load_workbook(calc_path, data_only=True)
    tx = B.demo_transactions(B.YEAR)
    inc_cats = set(B.INCOME_CATS)
    by_month = collections.defaultdict(lambda: collections.defaultdict(float))
    for d, _, cat, amt in tx:
        by_month[d.month][cat] += amt

    # Transactions: auto type + month
    ws = wb[B.S_TX]
    for i, (d, _, cat, _) in enumerate(tx):
        r = B.TX_FIRST + i
        check(ws[f"F{r}"].value == ("Income" if cat in inc_cats else "Expense"), f"Tx type row {r}")
        check(ws[f"G{r}"].value == d.month, f"Tx month row {r}")

    # Monthly dashboard (January is selected)
    ws = wb[B.S_MONTH]
    jan = by_month[1]
    inc = sum(v for c, v in jan.items() if c in inc_cats)
    exp = sum(v for c, v in jan.items() if c not in inc_cats)
    check(close(ws["B9"].value, inc), f"Monthly income {ws['B9'].value} vs {inc}")
    check(close(ws["C9"].value, exp), f"Monthly spent {ws['C9'].value} vs {exp}")
    check(close(ws["E9"].value, inc - exp), "Monthly left over")
    savings = sum(v for c, v in jan.items() if dict(B.EXPENSE_CATS).get(c) == "Savings")
    check(close(ws["G9"].value, (savings + inc - exp) / inc, 0.0001), f"Savings rate {ws['G9'].value}")
    found = 0
    for row in ws.iter_rows(min_row=14, max_row=80, min_col=2, max_col=5):
        name = row[0].value
        if name in jan and name not in inc_cats:
            check(close(row[3].value, jan[name]), f"Monthly actual {name}: {row[3].value} vs {jan[name]}")
            found += 1
    check(found >= 15, f"Monthly: matched {found} expense categories")

    # Annual overview: monthly totals and running balance
    ws = wb[B.S_ANNUAL]
    tot_rows = {ws[f"B{r}"].value: r for r in range(7, 70) if isinstance(ws[f"B{r}"].value, str)}
    ri, re_ = tot_rows["Total income"], tot_rows["Total expenses"]
    run = 0
    for m in range(1, 13):
        col = B.get_column_letter(2 + m)
        inc_m = sum(v for c, v in by_month[m].items() if c in inc_cats)
        exp_m = sum(v for c, v in by_month[m].items() if c not in inc_cats)
        run += inc_m - exp_m
        check(close(ws[f"{col}{ri}"].value, inc_m), f"Annual income {col}")
        check(close(ws[f"{col}{re_}"].value, exp_m), f"Annual expenses {col}")
        check(close(ws[f"{col}{tot_rows['Running balance']}"].value, run), f"Running balance {col}")
    print(f"  annual income {ws[f'O{ri}'].value:,.2f} | expenses {ws[f'O{re_}'].value:,.2f}")

    verify_debt(wb, B.DEMO_DEBTS, 200, "Snowball")
    verify_holiday(wb)


def verify_debt(wb, debts, extra, strategy):
    ws = wb[B.S_DEBT]
    expected = simulate_debts(debts, extra, strategy)
    for r in range(B.DEBT_FIRST, B.DEBT_FIRST + len(debts)):
        name = ws[f"B{r}"].value
        months, paid = expected[name]
        check(ws[f"G{r}"].value == months, f"Debt {name} payoff month {ws[f'G{r}'].value} vs {months}")
        check(close(ws[f"J{r}"].value, paid, 0.05), f"Debt {name} total paid {ws[f'J{r}'].value} vs {paid}")
        print(f"  {name:<14} priority {ws[f'F{r}'].value}  payoff month {ws[f'G{r}'].value}  "
              f"interest {ws[f'I{r}'].value:,.2f}")
    print(f"  debt-free: {ws['F6'].value}  months: {ws['H6'].value}  total interest: {ws['J6'].value:,.2f}")


def verify_holiday(wb):
    ws = wb[B.S_HOL]
    spent_g = sum(g[5] for g in B.DEMO_GIFTS)
    spent_o = sum(a for _, a in B.DEMO_HOLIDAY_OTHER.values())
    check(close(ws["G6"].value, spent_g + spent_o), f"Holiday spent {ws['G6'].value} vs {spent_g + spent_o}")
    check(close(ws["I6"].value, 1500 - spent_g - spent_o), "Holiday left in budget")
    bought = sum(1 for g in B.DEMO_GIFTS if g[6] in ("Bought", "Wrapped", "Given"))
    check(close(ws["E6"].value, bought / len(B.DEMO_GIFTS), 0.0001), "Holiday % bought")
    print(f"  holiday spent {ws['G6'].value:,.2f}, left {ws['I6'].value:,.2f}, bought {ws['E6'].value:.0%}")


def main():
    srcs = sorted(glob.glob(os.path.join(DIST, "*.xlsx")))
    if not srcs:
        sys.exit("No workbooks in dist/ — run build_workbook.py first.")
    with tempfile.TemporaryDirectory() as tmp:
        # Extra scenario: the same debts with the Avalanche strategy and a different extra payment
        variant_dir = os.path.join(tmp, "variant")
        os.makedirs(variant_dir)
        variant = os.path.join(variant_dir, "Debt-Avalanche.xlsx")
        vb = load_workbook(os.path.join(DIST, "Debt-Payoff-Planner-DEMO.xlsx"))
        vb[B.S_DEBT]["C6"] = "Avalanche"
        vb[B.S_DEBT]["C7"] = 75
        vb.save(variant)
        recalc([variant], tmp)
        print("Debt-Payoff-Planner (Avalanche, extra 75)")
        verify_debt(load_workbook(os.path.join(tmp, "Debt-Avalanche.xlsx"), data_only=True),
                    B.DEMO_DEBTS, 75, "Avalanche")

        recalc(srcs, tmp)
        for src in srcs:
            calc = os.path.join(tmp, os.path.basename(src))
            print(os.path.basename(src))
            n = scan_errors(src, calc)
            print(f"  {n:,} formulas scanned")
            base = os.path.basename(src)
            if base == "2027-Money-Planner-DEMO.xlsx":
                verify_bundle_demo(calc)
            elif base == "Debt-Payoff-Planner-DEMO.xlsx":
                verify_debt(load_workbook(calc, data_only=True), B.DEMO_DEBTS, 200, "Snowball")
            elif base == "Holiday-Budget-Gift-Tracker-DEMO.xlsx":
                verify_holiday(load_workbook(calc, data_only=True))
    print("\nALL CHECKS PASSED" if not failures else f"\n{len(failures)} CHECK(S) FAILED")
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
