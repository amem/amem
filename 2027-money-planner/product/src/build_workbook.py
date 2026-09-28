"""Build the 2027 Money Planner spreadsheets (.xlsx, Excel + Google Sheets).

Usage:
    python build_workbook.py            # builds every edition into ../dist
    python build_workbook.py --outdir X # custom output directory

Editions (one engine, several listings):
    bundle   - Complete 2027 Money Planner (all tabs)
    holiday  - Holiday Budget & Gift Tracker (standalone)
    debt     - Debt Snowball & Avalanche Payoff Planner (standalone)
Each edition is written twice: a clean copy (what the buyer uses) and a
"DEMO" copy filled with sample data (for mockups and for buyers to explore).

Only formulas that exist in both Excel and Google Sheets are used
(SUMIFS, COUNTIF, INDEX, MATCH, RANK, REPT, DATE, ...). No macros.
"""

import argparse
import datetime as dt
import os
import random

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# --------------------------------------------------------------------------
# Brand + style
# --------------------------------------------------------------------------
BRAND = "AMEM Sheets"
PRODUCT = "2027 Money Planner"
YEAR = 2027

FONT = "Arial"
C_PRIMARY = "2F5D50"   # deep green
C_SAGE = "8FB9A8"
C_MINT = "EEF4F1"
C_CREAM = "FFF8E7"     # input cells
C_GOLD = "E9B872"
C_RED = "B23A48"
C_TEXT = "2B2B2B"
C_MUTED = "7A8580"
C_LINE = "D9E2DE"
C_WHITE = "FFFFFF"

MONEY = '#,##0.00;[Red]-#,##0.00'
MONEY0 = '#,##0;[Red]-#,##0'
PCT = '0%'
DATE_FMT = 'dd mmm yyyy'
MONTH_FMT = 'mmm yyyy'

MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]
MON3 = [m[:3] for m in MONTHS]

thin = Side(style="thin", color=C_LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL_PRIMARY = PatternFill("solid", fgColor=C_PRIMARY)
FILL_MINT = PatternFill("solid", fgColor=C_MINT)
FILL_INPUT = PatternFill("solid", fgColor=C_CREAM)
FILL_SAGE = PatternFill("solid", fgColor=C_SAGE)
FILL_WHITE = PatternFill("solid", fgColor=C_WHITE)
FILL_RED_LIGHT = PatternFill("solid", fgColor="F6D5D9")
FILL_GREEN_LIGHT = PatternFill("solid", fgColor="D6EDE2")
FILL_GOLD_LIGHT = PatternFill("solid", fgColor="FBEBD0")
FILL_BLUE_LIGHT = PatternFill("solid", fgColor="DCE8F5")
FILL_GREY_LIGHT = PatternFill("solid", fgColor="EDEDED")

# Rows / ranges shared between sheets (kept in one place so formulas agree)
INC_FIRST, INC_LAST = 6, 15           # Categories: income rows
EXP_FIRST, EXP_LAST = 20, 49          # Categories: expense rows
TX_FIRST, TX_LAST = 6, 1005           # Transactions rows
DEBT_FIRST, DEBT_LAST = 12, 17        # Debt Payoff: 6 debts
SCHED_START_ROW = 7                   # Debt Schedule: "month 0" row
SCHED_MONTHS = 360                    # 30 years
SCHED_FIRST = SCHED_START_ROW + 1
SCHED_LAST = SCHED_START_ROW + SCHED_MONTHS

# Sheet names
S_START = "Start Here"
S_CAT = "Categories"
S_TX = "Transactions"
S_MONTH = "Monthly Dashboard"
S_ANNUAL = "Annual Overview"
S_BILLS = "Bills"
S_SAVE = "Savings Goals"
S_NW = "Net Worth"
S_DEBT = "Debt Payoff"
S_SCHED = "Debt Schedule"
S_HOL = "Holiday Planner"
S_LISTS = "Lists"


def q(name):
    """Quote a sheet name for use inside a formula."""
    return "'" + name.replace("'", "''") + "'"


TX = q(S_TX)
CAT = q(S_CAT)
START = q(S_START)
LISTS = q(S_LISTS)
DEBT = q(S_DEBT)
SCHED = q(S_SCHED)

TX_DATE = f"{TX}!$B${TX_FIRST}:$B${TX_LAST}"
TX_CAT = f"{TX}!$D${TX_FIRST}:$D${TX_LAST}"
TX_AMT = f"{TX}!$E${TX_FIRST}:$E${TX_LAST}"
TX_TYPE = f"{TX}!$F${TX_FIRST}:$F${TX_LAST}"
TX_MON = f"{TX}!$G${TX_FIRST}:$G${TX_LAST}"
CAT_INC = f"{CAT}!$B${INC_FIRST}:$B${INC_LAST}"
CAT_EXP = f"{CAT}!$B${EXP_FIRST}:$B${EXP_LAST}"
YEAR_CELL = f"{START}!$D$9"
START_BAL_CELL = f"{START}!$D$10"


def progress_bar(pct_expr, width=10):
    """Text progress bar that renders identically in Excel and Google Sheets."""
    filled = f"ROUND(MAX(0,MIN(1,{pct_expr}))*{width},0)"
    return f'REPT("■",{filled})&REPT("□",{width}-{filled})'


# --------------------------------------------------------------------------
# Low-level cell helpers
# --------------------------------------------------------------------------
def font(size=10, bold=False, color=C_TEXT, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def put(ws, ref, value=None, *, size=10, bold=False, color=C_TEXT, italic=False,
        fill=None, fmt=None, align=None, wrap=False, border=False, valign="center"):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font(size, bold, color, italic)
    if fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
    if border:
        c.border = BORDER
    return c


def header_row(ws, row, col_start, labels, fill=FILL_PRIMARY, color=C_WHITE, height=22):
    for i, label in enumerate(labels):
        ref = f"{get_column_letter(col_start + i)}{row}"
        put(ws, ref, label, bold=True, color=color, fill=fill, align="center",
            border=True, wrap=True)
    ws.row_dimensions[row].height = height


def title_block(ws, title, subtitle, width_cols=10):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2.5
    put(ws, "B2", title, size=20, bold=True, color=C_PRIMARY)
    put(ws, "B3", subtitle, size=10, italic=True, color=C_MUTED)
    ws.row_dimensions[2].height = 30
    # thin accent line under the title
    for col in range(2, 2 + width_cols):
        ws.cell(row=4, column=col).border = Border(top=Side(style="medium", color=C_SAGE))


def input_cell(ws, ref, value=None, fmt=None, align=None):
    return put(ws, ref, value, fill=FILL_INPUT, fmt=fmt, border=True, align=align)


def calc_cell(ws, ref, formula, fmt=None, bold=False, fill=None, align=None, color=C_TEXT):
    return put(ws, ref, formula, fmt=fmt, bold=bold, fill=fill, border=True,
               align=align, color=color)


def set_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def add_list_validation(ws, rng, source, prompt=None, strict=True):
    dv = DataValidation(type="list", formula1=source, allow_blank=True)
    dv.showDropDown = False  # (openpyxl quirk: False == show the arrow)
    if strict:
        dv.errorStyle = "stop"
        dv.error = "Please pick a value from the list."
        dv.errorTitle = "Invalid entry"
        dv.showErrorMessage = True
    else:
        dv.errorStyle = "warning"
        dv.showErrorMessage = True
        dv.error = "This value is not in the list. Keep it anyway?"
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv


def print_setup(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def kpi_card(ws, col, row, label, formula, fmt=MONEY, sub=None, fill=FILL_MINT, width_cols=2):
    """Label / big value / sub-caption spanning `width_cols` merged columns."""
    c1 = get_column_letter(col)
    c2 = get_column_letter(col + width_cols - 1)
    for r in (row, row + 1, row + 2):
        if width_cols > 1:
            ws.merge_cells(f"{c1}{r}:{c2}{r}")
        for cc in range(col, col + width_cols):
            ws.cell(row=r, column=cc).fill = fill
    put(ws, f"{c1}{row}", label, size=9, bold=True, color=C_MUTED, fill=fill, align="center")
    put(ws, f"{c1}{row + 1}", formula, size=16, bold=True, color=C_PRIMARY, fill=fill,
        fmt=fmt, align="center")
    put(ws, f"{c1}{row + 2}", sub if sub is not None else "", size=8, italic=True,
        color=C_MUTED, fill=fill, align="center")
    ws.row_dimensions[row + 1].height = 26


def style_chart(chart, title, width=16, height=8):
    chart.title = title
    chart.width = width
    chart.height = height
    chart.legend.position = "b"
    return chart


def color_series(chart, colors):
    for s, colr in zip(chart.series, colors):
        s.graphicalProperties.solidFill = colr
        s.graphicalProperties.line.solidFill = colr


# --------------------------------------------------------------------------
# Default + demo content
# --------------------------------------------------------------------------
INCOME_CATS = ["Salary", "Side Hustle", "Other Income"]
EXPENSE_CATS = [
    ("Rent / Mortgage", "Needs"), ("Utilities", "Needs"), ("Groceries", "Needs"),
    ("Transportation", "Needs"), ("Insurance", "Needs"), ("Phone & Internet", "Needs"),
    ("Health", "Needs"), ("Childcare / Education", "Needs"),
    ("Dining Out", "Wants"), ("Entertainment", "Wants"), ("Shopping", "Wants"),
    ("Subscriptions", "Wants"), ("Personal Care", "Wants"), ("Gifts & Holidays", "Wants"),
    ("Travel", "Wants"),
    ("Emergency Fund", "Savings"), ("Retirement", "Savings"), ("Sinking Funds", "Savings"),
    ("Credit Card Payment", "Debt"), ("Loan Payment", "Debt"),
]
DEMO_INCOME_PLAN = {"Salary": 4200, "Side Hustle": 400, "Other Income": 50}
DEMO_EXPENSE_BUDGET = {
    "Rent / Mortgage": 1400, "Utilities": 180, "Groceries": 480, "Transportation": 220,
    "Insurance": 160, "Phone & Internet": 90, "Health": 60, "Childcare / Education": 0,
    "Dining Out": 180, "Entertainment": 80, "Shopping": 150, "Subscriptions": 45,
    "Personal Care": 50, "Gifts & Holidays": 60, "Travel": 100,
    "Emergency Fund": 300, "Retirement": 250, "Sinking Funds": 100,
    "Credit Card Payment": 250, "Loan Payment": 200,
}


def demo_transactions(year):
    """Deterministic, realistic sample year of transactions."""
    rnd = random.Random(2027)
    rows = []

    def add(m, d, desc, cat, amt):
        rows.append((dt.date(year, m, d), desc, cat, round(amt, 2)))

    for m in range(1, 13):
        add(m, 1, "Paycheck", "Salary", 2100)
        add(m, 15, "Paycheck", "Salary", 2100)
        add(m, rnd.randint(8, 25), "Freelance design job", "Side Hustle", rnd.choice([250, 320, 400, 480, 550]))
        if m in (3, 7, 11):
            add(m, 20, "Cashback / interest", "Other Income", rnd.choice([35, 60, 85]))
        add(m, 1, "Rent", "Rent / Mortgage", 1400)
        add(m, 5, "Electricity & water", "Utilities", rnd.uniform(140, 215))
        for d in (3, 10, 17, 24):
            add(m, d, "Supermarket", "Groceries", rnd.uniform(92, 138))
        for d in rnd.sample(range(2, 28), 3):
            add(m, d, rnd.choice(["Fuel", "Metro card", "Car wash", "Parking"]), "Transportation", rnd.uniform(35, 85))
        add(m, 7, "Car & home insurance", "Insurance", 160)
        add(m, 9, "Phone + fibre", "Phone & Internet", 89.99)
        if rnd.random() < 0.55:
            add(m, rnd.randint(2, 27), "Pharmacy", "Health", rnd.uniform(15, 70))
        for d in rnd.sample(range(1, 28), rnd.randint(3, 5)):
            add(m, d, rnd.choice(["Pizza night", "Coffee shop", "Brunch", "Sushi", "Takeaway"]), "Dining Out", rnd.uniform(14, 62))
        for d in rnd.sample(range(1, 28), 2):
            add(m, d, rnd.choice(["Cinema", "Concert", "Bowling", "Museum"]), "Entertainment", rnd.uniform(18, 55))
        for d in rnd.sample(range(1, 28), rnd.randint(1, 3)):
            add(m, d, rnd.choice(["Clothes", "Home decor", "Online order", "Shoes"]), "Shopping", rnd.uniform(25, 95))
        add(m, 12, "Streaming", "Subscriptions", 15.49)
        add(m, 12, "Music", "Subscriptions", 10.99)
        add(m, 18, "Cloud storage", "Subscriptions", 2.99)
        add(m, rnd.randint(5, 25), "Haircut", "Personal Care", rnd.uniform(25, 45))
        gifts = {11: 180, 12: 320, 2: 60, 5: 70}.get(m, 0)
        if gifts or rnd.random() < 0.35:
            add(m, rnd.randint(5, 22), "Gifts", "Gifts & Holidays", gifts or rnd.uniform(20, 45))
        if m in (4, 8):
            add(m, 14, "Weekend trip", "Travel", rnd.uniform(380, 520))
        add(m, 16, "Transfer to emergency fund", "Emergency Fund", 300)
        add(m, 16, "Retirement contribution", "Retirement", 250)
        add(m, 16, "Car repair + gifts fund", "Sinking Funds", 100)
        add(m, 20, "Credit card payment", "Credit Card Payment", 250)
        add(m, 22, "Student loan", "Loan Payment", 200)
    rows.sort(key=lambda r: r[0])
    return rows


DEMO_BILLS = [
    ("Rent", 1, 1400, "No"), ("Electricity & water", 5, 180, "Yes"), ("Car & home insurance", 7, 160, "Yes"),
    ("Phone + fibre", 9, 89.99, "Yes"), ("Streaming", 12, 15.49, "Yes"), ("Music", 12, 10.99, "Yes"),
    ("Cloud storage", 18, 2.99, "Yes"), ("Credit card", 20, 250, "No"), ("Student loan", 22, 200, "Yes"),
]
DEMO_GOALS = [
    ("Emergency fund (3 months)", 9000, 3600, dt.date(2027, 12, 31)),
    ("New laptop", 1400, 650, dt.date(2027, 6, 30)),
    ("Summer vacation", 2500, 900, dt.date(2027, 7, 15)),
    ("Car down payment", 5000, 1200, dt.date(2028, 3, 31)),
    ("Holiday gifts 2027", 800, 150, dt.date(2027, 11, 30)),
]
DEMO_DEBTS = [
    ("Store card", 850, 26.99, 35), ("Credit card", 4200, 21.49, 110),
    ("Car loan", 9800, 6.9, 245), ("Student loan", 14500, 4.5, 160),
]
DEMO_ASSETS = [("Checking account", 2800, 60), ("Emergency fund", 3600, 300),
               ("Retirement account", 18500, 380), ("Car value", 11000, -90)]
DEMO_LIABS = [("Credit cards", 5050, -180), ("Car loan", 9800, -190), ("Student loan", 14500, -110)]
DEMO_GIFTS = [
    ("Mom", "Family", 80, "Cashmere scarf", "Department store", 72, "Wrapped"),
    ("Dad", "Family", 80, "Coffee grinder", "Online", 64.5, "Bought"),
    ("Sara (sister)", "Family", 60, "Book set + candle", "Bookshop", 58, "Bought"),
    ("Adam (brother)", "Family", 60, "Headphones", "Online", 0, "Ordered"),
    ("Grandma", "Family", 40, "Photo calendar", "Print shop", 0, "Idea"),
    ("Lina (niece)", "Kids", 35, "Art kit", "Toy store", 32, "Wrapped"),
    ("Yusuf (nephew)", "Kids", 35, "LEGO set", "Toy store", 39.99, "Bought"),
    ("Nora", "Friends", 30, "Plant + pot", "Garden center", 0, "Idea"),
    ("Karim", "Friends", 30, "Board game", "Online", 27, "Given"),
    ("Office Secret Santa", "Work", 25, "Gourmet chocolate box", "Online", 24, "Bought"),
    ("Neighbors", "Other", 20, "Homemade cookies", "Home", 0, "Idea"),
    ("Teacher", "Other", 20, "Gift card", "Online", 20, "Given"),
]
HOLIDAY_OTHER = ["Food & dinner", "Decorations", "Wrapping & cards", "Travel",
                 "Parties & events", "Outfits", "Charity & donations", "Stocking stuffers", "Other"]
DEMO_HOLIDAY_OTHER = {"Food & dinner": (250, 138.4), "Decorations": (80, 64.2), "Wrapping & cards": (40, 22.5),
                      "Travel": (200, 0), "Parties & events": (100, 45), "Outfits": (80, 0),
                      "Charity & donations": (50, 50), "Stocking stuffers": (40, 12), "Other": (30, 0)}


# --------------------------------------------------------------------------
# Sheets
# --------------------------------------------------------------------------
def build_lists(wb):
    ws = wb.create_sheet(S_LISTS)
    cols = {
        "A": ["Months"] + MONTHS,
        "B": ["Groups", "Needs", "Wants", "Savings", "Debt"],
        "C": ["Bill status", "Paid", "Unpaid"],
        "D": ["Gift status", "Idea", "Ordered", "Bought", "Wrapped", "Given"],
        "E": ["Gift group", "Family", "Friends", "Kids", "Work", "Other"],
        "F": ["Strategy", "Snowball", "Avalanche"],
        "G": ["Yes/No", "Yes", "No"],
    }
    for col, vals in cols.items():
        for i, v in enumerate(vals, start=1):
            ws[f"{col}{i}"] = v
            ws[f"{col}{i}"].font = font(bold=(i == 1))
    ws.sheet_state = "hidden"
    return ws


def build_start(wb, edition, tabs):
    ws = wb.active
    ws.title = S_START
    names = {
        "bundle": (PRODUCT, "Budget • Bills • Savings • Debt payoff • Net worth • Holiday gifts — Excel & Google Sheets"),
        "holiday": ("Holiday Budget & Gift Tracker", "Plan every gift, stay on budget, and enjoy the holidays — Excel & Google Sheets"),
        "debt": ("Debt Payoff Planner", "Snowball or Avalanche — see your exact debt-free date — Excel & Google Sheets"),
    }
    title, subtitle = names[edition]
    title_block(ws, title, subtitle, width_cols=6)
    set_widths(ws, {"B": 4, "C": 28, "D": 22, "E": 60, "F": 4, "G": 4})

    row = 6
    put(ws, f"B{row}", "WELCOME", size=12, bold=True, color=C_PRIMARY)
    row += 1
    put(ws, f"C{row}", "Thank you for your purchase! Cream-colored cells are for you to type in; "
                        "everything else calculates automatically.", wrap=True, italic=True, color=C_MUTED)
    ws.merge_cells(f"C{row}:E{row}")
    ws.row_dimensions[row].height = 28

    if "budget" in tabs:
        row = 9
        put(ws, f"C{row}", "Planner year", bold=True, border=True, fill=FILL_MINT)
        input_cell(ws, f"D{row}", YEAR, fmt="0", align="center")
        put(ws, f"E{row}", "Change this to reuse the planner for any year (e.g. 2026 to start today).",
            italic=True, color=C_MUTED, size=9)
        row = 10
        put(ws, f"C{row}", "Starting balance (Jan 1)", bold=True, border=True, fill=FILL_MINT)
        input_cell(ws, f"D{row}", 0, fmt=MONEY)
        put(ws, f"E{row}", "Money in your main account on day one (used for the running balance).",
            italic=True, color=C_MUTED, size=9)
        row = 12
    else:
        row = 9

    put(ws, f"B{row}", "QUICK START", size=12, bold=True, color=C_PRIMARY)
    row += 1
    steps = {
        "bundle": [
            "Set your planner year and starting balance above.",
            "Go to Categories: rename income & expense categories and type your monthly budget for each.",
            "Log every income and expense in Transactions (date, description, category, amount).",
            "Open Monthly Dashboard and pick a month to see budget vs. actual instantly.",
            "Annual Overview, Bills, Savings Goals, Net Worth and Debt Payoff update as you go.",
            "Holiday Planner: list who you are buying for and track every gift against your budget.",
        ],
        "holiday": [
            "Open Holiday Planner and set your holiday date and total budget.",
            "List everyone you are buying for, their group and gift budget.",
            "Add gift ideas and change the status as you go: Idea → Ordered → Bought → Wrapped → Given.",
            "Track food, decor, travel and other holiday costs in the 'Other holiday costs' table.",
        ],
        "debt": [
            "Open Debt Payoff and list up to 6 debts with balance, APR (e.g. 19.99) and minimum payment.",
            "Choose Snowball (smallest balance first) or Avalanche (highest interest first).",
            "Add any extra amount you can pay each month and your first payment month.",
            "See each debt's payoff date, total interest and your debt-free date. Full plan in Debt Schedule.",
        ],
    }[edition]
    for i, s in enumerate(steps, start=1):
        put(ws, f"B{row}", i, bold=True, color=C_WHITE, fill=FILL_PRIMARY, align="center")
        put(ws, f"C{row}", s, wrap=True)
        ws.merge_cells(f"C{row}:E{row}")
        ws.row_dimensions[row].height = 20
        row += 1

    row += 1
    put(ws, f"B{row}", "WHAT'S INSIDE", size=12, bold=True, color=C_PRIMARY)
    row += 1
    descriptions = {
        S_CAT: "Your income & expense categories, monthly budgets and a 50/30/20 check.",
        S_TX: "One simple log for every income and expense (1,000 rows).",
        S_MONTH: "Pick a month: budget vs. actual, left over, savings rate, charts.",
        S_ANNUAL: "All 12 months side by side with totals, averages and running balance.",
        S_BILLS: "Recurring bills with due dates and a paid/unpaid check for every month.",
        S_SAVE: "Savings goals with progress bars and how much to save per month.",
        S_NW: "Track assets and debts monthly and watch your net worth grow.",
        S_DEBT: "Snowball / Avalanche calculator with payoff dates and total interest.",
        S_SCHED: "Month-by-month payoff schedule for every debt (auto-calculated).",
        S_HOL: "Gift list, status tracker, budget per person and other holiday costs.",
    }
    for name in [n for n in descriptions if n in wb.sheetnames or n in PLANNED_SHEETS]:
        c = put(ws, f"C{row}", name, bold=True, color=C_PRIMARY, border=True, fill=FILL_MINT)
        c.hyperlink = f"#{q(name)}!A1"
        put(ws, f"D{row}", descriptions[name], wrap=True, border=True)
        ws.merge_cells(f"D{row}:E{row}")
        row += 1

    row += 1
    put(ws, f"B{row}", "COLOR KEY", size=12, bold=True, color=C_PRIMARY)
    row += 1
    put(ws, f"C{row}", "Type here", border=True, fill=FILL_INPUT, align="center")
    put(ws, f"D{row}", "Your inputs", italic=True, color=C_MUTED)
    row += 1
    put(ws, f"C{row}", "Auto", border=True, align="center")
    put(ws, f"D{row}", "Formulas — no need to edit", italic=True, color=C_MUTED)
    row += 1
    put(ws, f"C{row}", "Total", border=True, fill=FILL_MINT, bold=True, align="center")
    put(ws, f"D{row}", "Totals & summaries", italic=True, color=C_MUTED)

    row += 2
    put(ws, f"B{row}", "GOOGLE SHEETS", size=12, bold=True, color=C_PRIMARY)
    row += 1
    put(ws, f"C{row}", "Upload this file to Google Drive → right-click → Open with → Google Sheets → "
                        "File → Save as Google Sheets. Works in any currency.", wrap=True)
    ws.merge_cells(f"C{row}:E{row}")
    ws.row_dimensions[row].height = 30
    row += 2
    put(ws, f"C{row}", f"© {BRAND}. For personal use only — please do not share or resell.",
        size=8, italic=True, color=C_MUTED)
    print_setup(ws, landscape=False)
    return ws


PLANNED_SHEETS = set()


def build_categories(wb, demo):
    ws = wb.create_sheet(S_CAT)
    title_block(ws, "Categories & Budget", "Rename categories and set a monthly budget. Fill lists from the top down.", 9)
    set_widths(ws, {"B": 26, "C": 16, "D": 16, "E": 3, "F": 14, "G": 16, "H": 14, "I": 16, "K": 26})

    put(ws, "B4", "INCOME", size=12, bold=True, color=C_PRIMARY)
    header_row(ws, 5, 2, ["Income category", "Planned / month"])
    for i, r in enumerate(range(INC_FIRST, INC_LAST + 1)):
        name = INCOME_CATS[i] if i < len(INCOME_CATS) else None
        input_cell(ws, f"B{r}", name)
        input_cell(ws, f"C{r}", DEMO_INCOME_PLAN.get(name) if demo and name else None, fmt=MONEY)
    put(ws, f"B{INC_LAST + 1}", "Total planned income", bold=True, fill=FILL_MINT, border=True)
    calc_cell(ws, f"C{INC_LAST + 1}", f"=SUM(C{INC_FIRST}:C{INC_LAST})", fmt=MONEY, bold=True, fill=FILL_MINT)

    put(ws, f"B{EXP_FIRST - 2}", "EXPENSES", size=12, bold=True, color=C_PRIMARY)
    header_row(ws, EXP_FIRST - 1, 2, ["Expense category", "Group", "Budget / month"])
    for i, r in enumerate(range(EXP_FIRST, EXP_LAST + 1)):
        name, grp = EXPENSE_CATS[i] if i < len(EXPENSE_CATS) else (None, None)
        input_cell(ws, f"B{r}", name)
        input_cell(ws, f"C{r}", grp, align="center")
        input_cell(ws, f"D{r}", DEMO_EXPENSE_BUDGET.get(name) if demo and name else None, fmt=MONEY)
    add_list_validation(ws, f"C{EXP_FIRST}:C{EXP_LAST}", f"={LISTS}!$B$2:$B$5", "Needs, Wants, Savings or Debt")
    tot = EXP_LAST + 1
    put(ws, f"B{tot}", "Total monthly budget", bold=True, fill=FILL_MINT, border=True)
    put(ws, f"C{tot}", "", fill=FILL_MINT, border=True)
    calc_cell(ws, f"D{tot}", f"=SUM(D{EXP_FIRST}:D{EXP_LAST})", fmt=MONEY, bold=True, fill=FILL_MINT)
    put(ws, f"B{tot + 1}", "Unassigned (income − budget)", bold=True, border=True)
    put(ws, f"C{tot + 1}", "", border=True)
    calc_cell(ws, f"D{tot + 1}", f"=C{INC_LAST + 1}-D{tot}", fmt=MONEY, bold=True)

    # 50/30/20 check
    put(ws, "F4", "50 / 30 / 20 CHECK", size=12, bold=True, color=C_PRIMARY)
    header_row(ws, 5, 6, ["Group", "Budget / month", "% of income", "Guideline"])
    guide = {"Needs": "≤ 50%", "Wants": "≤ 30%", "Savings": "20%+ (with Debt)", "Debt": ""}
    for i, g in enumerate(["Needs", "Wants", "Savings", "Debt"]):
        r = 6 + i
        put(ws, f"F{r}", g, bold=True, border=True, fill=FILL_MINT)
        calc_cell(ws, f"G{r}", f"=SUMIF($C${EXP_FIRST}:$C${EXP_LAST},F{r},$D${EXP_FIRST}:$D${EXP_LAST})", fmt=MONEY)
        calc_cell(ws, f"H{r}", f"=IFERROR(G{r}/$C${INC_LAST + 1},0)", fmt=PCT, align="center")
        put(ws, f"I{r}", guide[g], border=True, align="center", color=C_MUTED, italic=True)
    put(ws, "F11", "Tip: the 50/30/20 rule is a starting point — adjust it to your life.",
        size=8, italic=True, color=C_MUTED)

    # Hidden helper: compact list of every category for the Transactions dropdown
    put(ws, "K5", "Dropdown list (auto)", bold=True)
    n_inc = f"COUNTA($B${INC_FIRST}:$B${INC_LAST})"
    for n in range(1, 41):
        r = 5 + n
        ws[f"K{r}"] = (f'=IF({n}<={n_inc},INDEX($B${INC_FIRST}:$B${INC_LAST},{n}),'
                       f'IF({n}-{n_inc}<=COUNTA($B${EXP_FIRST}:$B${EXP_LAST}),'
                       f'INDEX($B${EXP_FIRST}:$B${EXP_LAST},{n}-{n_inc}),""))')
        ws[f"K{r}"].font = font(color=C_MUTED)
    ws.column_dimensions["K"].hidden = True
    ws.freeze_panes = "A5"
    print_setup(ws)
    return ws


def build_transactions(wb, demo):
    ws = wb.create_sheet(S_TX)
    title_block(ws, "Transactions", "Log every income and expense here. Type and month fill in automatically.", 7)
    set_widths(ws, {"B": 14, "C": 32, "D": 24, "E": 14, "F": 14, "G": 12, "H": 28})
    header_row(ws, 5, 2, ["Date", "Description", "Category", "Amount", "Type (auto)", "Month (auto)", "Notes"])
    rows = demo_transactions(YEAR) if demo else []
    for i, r in enumerate(range(TX_FIRST, TX_LAST + 1)):
        d = rows[i] if i < len(rows) else None
        input_cell(ws, f"B{r}", d[0] if d else None, fmt=DATE_FMT, align="center")
        input_cell(ws, f"C{r}", d[1] if d else None)
        input_cell(ws, f"D{r}", d[2] if d else None)
        input_cell(ws, f"E{r}", d[3] if d else None, fmt=MONEY)
        calc_cell(ws, f"F{r}",
                  f'=IF(D{r}="","",IF(COUNTIF({CAT_INC},D{r})>0,"Income",'
                  f'IF(COUNTIF({CAT_EXP},D{r})>0,"Expense","Check category")))', align="center")
        calc_cell(ws, f"G{r}",
                  f'=IF(B{r}="","",IFERROR(IF(YEAR(B{r})={YEAR_CELL},MONTH(B{r}),"Other year"),"Check date"))',
                  align="center")
        input_cell(ws, f"H{r}")
    add_list_validation(ws, f"D{TX_FIRST}:D{TX_LAST}", f"={CAT}!$K$6:$K$45",
                        "Pick a category (edit the list on the Categories tab)", strict=False)
    rng = f"F{TX_FIRST}:F{TX_LAST}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Income"'],
                                                  font=Font(name=FONT, color="2E7D4F", bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Check category"'],
                                                  fill=FILL_RED_LIGHT, font=Font(name=FONT, color=C_RED, bold=True)))
    ws.conditional_formatting.add(f"G{TX_FIRST}:G{TX_LAST}",
                                  FormulaRule(formula=[f'ISTEXT(G{TX_FIRST})'], fill=FILL_GOLD_LIGHT))
    ws.freeze_panes = "B6"
    ws.auto_filter.ref = f"B5:H{TX_LAST}"
    print_setup(ws)
    return ws


def build_monthly(wb):
    ws = wb.create_sheet(S_MONTH)
    title_block(ws, "Monthly Dashboard", "Choose a month — everything below updates automatically.", 12)
    set_widths(ws, {"B": 24, "C": 13, "D": 14, "E": 14, "F": 14, "G": 10, "H": 14,
                    "I": 3, "J": 15, "K": 15, "L": 15, "M": 15})
    put(ws, "B6", "Select month ▸", bold=True, color=C_PRIMARY, size=11)
    input_cell(ws, "C6", "January", align="center")
    ws["C6"].font = font(12, True, C_PRIMARY)
    ws.merge_cells("C6:D6")
    add_list_validation(ws, "C6", f"={LISTS}!$A$2:$A$13", "Pick a month")
    calc_cell(ws, "E6", f"=MATCH(C6,{LISTS}!$A$2:$A$13,0)", fmt="0", align="center", color=C_MUTED)
    put(ws, "F6", "← month #", size=8, italic=True, color=C_MUTED)
    m = "$E$6"

    # Layout
    inc_label = 12
    inc_first = inc_label + 2
    inc_last = inc_first + (INC_LAST - INC_FIRST)
    inc_tot = inc_last + 1
    exp_label = inc_tot + 2
    exp_first = exp_label + 2
    exp_last = exp_first + (EXP_LAST - EXP_FIRST)
    exp_tot = exp_last + 1
    grp_first = exp_first
    sav_row = grp_first + 2  # "Savings" row in the group table

    inc = f'SUMIFS({TX_AMT},{TX_TYPE},"Income",{TX_MON},{m})'
    exp = f'SUMIFS({TX_AMT},{TX_TYPE},"Expense",{TX_MON},{m})'
    kpi_card(ws, 2, 8, "INCOME", f"={inc}", width_cols=1)
    kpi_card(ws, 3, 8, "SPENT", f"={exp}", width_cols=2)
    kpi_card(ws, 5, 8, "LEFT OVER", "=B9-C9", width_cols=2)
    kpi_card(ws, 7, 8, "SAVINGS RATE", f"=IFERROR((L{sav_row}+E9)/B9,0)", fmt=PCT, width_cols=2)
    ws["B10"] = f'="of "&TEXT({CAT}!C{INC_LAST + 1},"#,##0")&" planned"'
    ws["C10"] = f'="of "&TEXT({CAT}!D{EXP_LAST + 1},"#,##0")&" budgeted"'
    ws["E10"] = "income − spent"
    ws["G10"] = "(savings + left over) ÷ income"
    for ref in ("B10", "C10", "E10", "G10"):
        ws[ref].font = font(8, italic=True, color=C_MUTED)
        ws[ref].alignment = Alignment(horizontal="center")
    ws.conditional_formatting.add("E9", CellIsRule(operator="lessThan", formula=["0"],
                                                   font=Font(name=FONT, color=C_RED, bold=True, size=16)))

    # Income table
    put(ws, f"B{inc_label}", "INCOME", size=12, bold=True, color=C_PRIMARY)
    header_row(ws, inc_label + 1, 2, ["Category", "Planned", "Actual", "Difference", "Progress"])
    ws.merge_cells(f"F{inc_label + 1}:G{inc_label + 1}")
    for i in range(INC_LAST - INC_FIRST + 1):
        r = inc_first + i
        src = INC_FIRST + i
        calc_cell(ws, f"B{r}", f'=IF({CAT}!B{src}="","",{CAT}!B{src})')
        calc_cell(ws, f"C{r}", f'=IF(B{r}="","",N({CAT}!C{src}))', fmt=MONEY)
        calc_cell(ws, f"D{r}", f'=IF(B{r}="","",SUMIFS({TX_AMT},{TX_CAT},B{r},{TX_MON},{m}))', fmt=MONEY)
        calc_cell(ws, f"E{r}", f'=IF(B{r}="","",D{r}-C{r})', fmt=MONEY)
        calc_cell(ws, f"F{r}", f'=IF(B{r}="","",{progress_bar(f"IFERROR(D{r}/C{r},0)")})', color=C_PRIMARY)
        ws[f"G{r}"].border = BORDER
        ws.merge_cells(f"F{r}:G{r}")
    put(ws, f"B{inc_tot}", "Total income", bold=True, fill=FILL_MINT, border=True)
    for col in "CDE":
        calc_cell(ws, f"{col}{inc_tot}", f"=SUM({col}{inc_first}:{col}{inc_last})", fmt=MONEY, bold=True, fill=FILL_MINT)
    put(ws, f"F{inc_tot}", "", fill=FILL_MINT, border=True)
    put(ws, f"G{inc_tot}", "", fill=FILL_MINT, border=True)
    ws.merge_cells(f"F{inc_tot}:G{inc_tot}")

    # Expense table
    put(ws, f"B{exp_label}", "EXPENSES", size=12, bold=True, color=C_PRIMARY)
    header_row(ws, exp_label + 1, 2, ["Category", "Group", "Budget", "Actual", "Remaining", "% used", "Progress"])
    for i in range(EXP_LAST - EXP_FIRST + 1):
        r = exp_first + i
        src = EXP_FIRST + i
        calc_cell(ws, f"B{r}", f'=IF({CAT}!B{src}="","",{CAT}!B{src})')
        calc_cell(ws, f"C{r}", f'=IF(B{r}="","",{CAT}!C{src})', align="center")
        calc_cell(ws, f"D{r}", f'=IF(B{r}="","",N({CAT}!D{src}))', fmt=MONEY)
        calc_cell(ws, f"E{r}", f'=IF(B{r}="","",SUMIFS({TX_AMT},{TX_CAT},B{r},{TX_MON},{m}))', fmt=MONEY)
        calc_cell(ws, f"F{r}", f'=IF(B{r}="","",D{r}-E{r})', fmt=MONEY)
        calc_cell(ws, f"G{r}", f'=IF(B{r}="","",IFERROR(E{r}/D{r},IF(E{r}>0,1,0)))', fmt=PCT, align="center")
        calc_cell(ws, f"H{r}", f'=IF(B{r}="","",{progress_bar(f"G{r}")})', color=C_PRIMARY)
    put(ws, f"B{exp_tot}", "Total expenses", bold=True, fill=FILL_MINT, border=True)
    put(ws, f"C{exp_tot}", "", fill=FILL_MINT, border=True)
    for col in "DEF":
        calc_cell(ws, f"{col}{exp_tot}", f"=SUM({col}{exp_first}:{col}{exp_last})", fmt=MONEY, bold=True, fill=FILL_MINT)
    calc_cell(ws, f"G{exp_tot}", f"=IFERROR(E{exp_tot}/D{exp_tot},0)", fmt=PCT, bold=True, fill=FILL_MINT, align="center")
    calc_cell(ws, f"H{exp_tot}", f'={progress_bar(f"G{exp_tot}")}', bold=True, fill=FILL_MINT, color=C_PRIMARY)
    ws.conditional_formatting.add(f"F{exp_first}:F{exp_tot}", CellIsRule(operator="lessThan", formula=["0"],
                                  fill=FILL_RED_LIGHT, font=Font(name=FONT, color=C_RED, bold=True)))
    ws.conditional_formatting.add(f"H{exp_first}:H{exp_tot}", FormulaRule(
        formula=[f'AND(ISNUMBER(G{exp_first}),G{exp_first}>1)'], font=Font(name=FONT, color=C_RED)))

    # Group summary + charts
    put(ws, f"J{exp_label}", "BY GROUP", size=12, bold=True, color=C_PRIMARY)
    header_row(ws, exp_label + 1, 10, ["Group", "Budget", "Actual", "Remaining"])
    for i, g in enumerate(["Needs", "Wants", "Savings", "Debt"]):
        r = grp_first + i
        put(ws, f"J{r}", g, bold=True, border=True, fill=FILL_MINT)
        calc_cell(ws, f"K{r}", f"=SUMIF($C${exp_first}:$C${exp_last},J{r},$D${exp_first}:$D${exp_last})", fmt=MONEY)
        calc_cell(ws, f"L{r}", f"=SUMIF($C${exp_first}:$C${exp_last},J{r},$E${exp_first}:$E${exp_last})", fmt=MONEY)
        calc_cell(ws, f"M{r}", f"=K{r}-L{r}", fmt=MONEY)
    ws.conditional_formatting.add(f"M{grp_first}:M{grp_first + 3}", CellIsRule(operator="lessThan", formula=["0"],
                                  font=Font(name=FONT, color=C_RED, bold=True)))

    bar = BarChart()
    bar.type = "col"
    bar.add_data(Reference(ws, min_col=11, max_col=12, min_row=exp_label + 1, max_row=grp_first + 3), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=10, min_row=grp_first, max_row=grp_first + 3))
    style_chart(bar, "Budget vs. actual by group", width=11.5, height=7.5)
    color_series(bar, [C_SAGE, C_PRIMARY])
    bar.y_axis.majorGridlines = None
    ws.add_chart(bar, f"J{grp_first + 5}")

    pie = PieChart()
    pie.add_data(Reference(ws, min_col=12, min_row=exp_label + 1, max_row=grp_first + 3), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=10, min_row=grp_first, max_row=grp_first + 3))
    style_chart(pie, "Where the money went", width=11.5, height=7.5)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showVal = False
    pie.dataLabels.showCatName = False
    pie.dataLabels.showSerName = False
    pie.dataLabels.showLegendKey = False
    for idx, colr in enumerate([C_PRIMARY, C_SAGE, C_GOLD, C_RED]):
        pt = DataPoint(idx=idx)
        pt.graphicalProperties.solidFill = colr
        pie.series[0].dPt.append(pt)
    ws.add_chart(pie, f"J{grp_first + 21}")

    ws.freeze_panes = "A7"
    print_setup(ws)
    return ws


def build_annual(wb):
    ws = wb.create_sheet(S_ANNUAL)
    title_block(ws, "Annual Overview", "Every month side by side. Totals, averages and your running balance.", 17)
    ws["B3"] = f'="Planner year "&{YEAR_CELL}&"  •  Everything here is automatic."'
    ws["B3"].font = font(10, italic=True, color=C_MUTED)
    set_widths(ws, {"B": 24, **{get_column_letter(c): 11 for c in range(3, 15)},
                    "O": 13, "P": 12, "Q": 13, "R": 13})
    # row 5: hidden month numbers, row 6: header
    for i in range(12):
        col = get_column_letter(3 + i)
        ws[f"{col}5"] = i + 1
        ws[f"{col}5"].font = font(7, color="BBBBBB")
        ws[f"{col}5"].alignment = Alignment(horizontal="center")
    header_row(ws, 6, 2, ["Category"] + MON3 + ["Total", "Avg / month", "Annual budget", "Variance"])
    put(ws, "R5", "(+ = better than plan)", size=7, italic=True, color=C_MUTED, align="center")

    # helper: months that already have transactions (row 4, small grey)
    put(ws, "B4", "months with data →", size=7, italic=True, color="BBBBBB", align="right")
    for i in range(12):
        col = get_column_letter(3 + i)
        ws[f"{col}4"] = f"=IF(COUNTIF({TX_MON},{col}$5)>0,1,0)"
        ws[f"{col}4"].font = font(7, color="BBBBBB")
        ws[f"{col}4"].alignment = Alignment(horizontal="center")
    ws["O4"] = "=MAX(1,SUM(C4:N4))"
    ws["O4"].font = font(7, color="BBBBBB")

    def section(label_row, first_src, last_src, kind):
        put(ws, f"B{label_row}", "INCOME" if kind == "inc" else "EXPENSES", bold=True, color=C_WHITE, fill=FILL_SAGE)
        for col in range(3, 19):
            ws.cell(row=label_row, column=col).fill = FILL_SAGE
        first = label_row + 1
        for i, src in enumerate(range(first_src, last_src + 1)):
            r = first + i
            calc_cell(ws, f"B{r}", f'=IF({CAT}!B{src}="","",{CAT}!B{src})')
            for mi in range(12):
                col = get_column_letter(3 + mi)
                calc_cell(ws, f"{col}{r}", f'=IF($B{r}="","",SUMIFS({TX_AMT},{TX_CAT},$B{r},{TX_MON},{col}$5))', fmt=MONEY0)
            calc_cell(ws, f"O{r}", f'=IF($B{r}="","",SUM(C{r}:N{r}))', fmt=MONEY0, bold=True)
            calc_cell(ws, f"P{r}", f'=IF($B{r}="","",O{r}/$O$4)', fmt=MONEY0)
            bud_col = "C" if kind == "inc" else "D"
            calc_cell(ws, f"Q{r}", f'=IF($B{r}="","",N({CAT}!{bud_col}{src})*12)', fmt=MONEY0)
            var = f"O{r}-Q{r}" if kind == "inc" else f"Q{r}-O{r}"
            calc_cell(ws, f"R{r}", f'=IF($B{r}="","",{var})', fmt=MONEY0)
        last = first + (last_src - first_src)
        tot = last + 1
        put(ws, f"B{tot}", "Total income" if kind == "inc" else "Total expenses", bold=True, fill=FILL_MINT, border=True)
        for col in range(3, 19):
            L = get_column_letter(col)
            calc_cell(ws, f"{L}{tot}", f"=SUM({L}{first}:{L}{last})", fmt=MONEY0, bold=True, fill=FILL_MINT)
        return first, last, tot

    _, _, inc_tot = section(7, INC_FIRST, INC_LAST, "inc")
    exp_label = inc_tot + 2
    _, exp_last, exp_tot = section(exp_label, EXP_FIRST, EXP_LAST, "exp")
    ws.conditional_formatting.add(f"R8:R{exp_tot}", CellIsRule(operator="lessThan", formula=["0"],
                                                               font=Font(name=FONT, color=C_RED, bold=True)))

    net = exp_tot + 2
    put(ws, f"B{net}", "Net (income − expenses)", bold=True, border=True, fill=FILL_GOLD_LIGHT)
    put(ws, f"B{net + 1}", "Savings rate", bold=True, border=True)
    put(ws, f"B{net + 2}", "Running balance", bold=True, border=True)
    for col in range(3, 16):
        L = get_column_letter(col)
        calc_cell(ws, f"{L}{net}", f"={L}{inc_tot}-{L}{exp_tot}", fmt=MONEY0, bold=True, fill=FILL_GOLD_LIGHT)
        calc_cell(ws, f"{L}{net + 1}", f"=IFERROR({L}{net}/{L}{inc_tot},0)", fmt=PCT, align="center")
    for mi in range(12):
        L = get_column_letter(3 + mi)
        calc_cell(ws, f"{L}{net + 2}", f"={START_BAL_CELL}+SUM($C{net}:{L}{net})", fmt=MONEY0)
    ws.conditional_formatting.add(f"C{net}:O{net}", CellIsRule(operator="lessThan", formula=["0"],
                                                               font=Font(name=FONT, color=C_RED, bold=True)))

    chart = BarChart()
    chart.type = "col"
    chart.add_data(Reference(ws, min_col=2, max_col=14, min_row=inc_tot), from_rows=True, titles_from_data=True)
    chart.add_data(Reference(ws, min_col=2, max_col=14, min_row=exp_tot), from_rows=True, titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=3, max_col=14, min_row=6))
    style_chart(chart, "Income vs. expenses by month", width=26, height=8)
    color_series(chart, [C_PRIMARY, C_GOLD])
    ws.add_chart(chart, f"B{net + 4}")
    ws.freeze_panes = "C7"
    print_setup(ws)
    return ws


def build_bills(wb, demo):
    ws = wb.create_sheet(S_BILLS)
    title_block(ws, "Bills Tracker", "List recurring bills once, then mark them Paid each month.", 17)
    set_widths(ws, {"B": 26, "C": 9, "D": 12, "E": 10, **{get_column_letter(c): 8 for c in range(6, 18)}})
    header_row(ws, 5, 2, ["Bill", "Due day", "Amount", "Autopay"] + MON3, height=24)
    first, last = 6, 30
    for i, r in enumerate(range(first, last + 1)):
        b = DEMO_BILLS[i] if demo and i < len(DEMO_BILLS) else None
        input_cell(ws, f"B{r}", b[0] if b else None)
        input_cell(ws, f"C{r}", b[1] if b else None, fmt="0", align="center")
        input_cell(ws, f"D{r}", b[2] if b else None, fmt=MONEY)
        input_cell(ws, f"E{r}", b[3] if b else None, align="center")
        for mi in range(12):
            L = get_column_letter(6 + mi)
            val = None
            if b:
                val = "Paid" if mi < 10 else ("Paid" if (mi == 10 and i % 3 != 0) else "Unpaid")
            input_cell(ws, f"{L}{r}", val, align="center")
    add_list_validation(ws, f"E{first}:E{last}", f"={LISTS}!$G$2:$G$3")
    add_list_validation(ws, f"F{first}:Q{last}", f"={LISTS}!$C$2:$C$3", "Paid or Unpaid")
    rng = f"F{first}:Q{last}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Paid"'], fill=FILL_GREEN_LIGHT,
                                                  font=Font(name=FONT, color="2E7D4F", bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Unpaid"'], fill=FILL_RED_LIGHT,
                                                  font=Font(name=FONT, color=C_RED)))
    t = last + 1
    put(ws, f"B{t}", "Total bills / month", bold=True, fill=FILL_MINT, border=True)
    put(ws, f"C{t}", "", fill=FILL_MINT, border=True)
    calc_cell(ws, f"D{t}", f"=SUM(D{first}:D{last})", fmt=MONEY, bold=True, fill=FILL_MINT)
    put(ws, f"E{t}", "Paid →", fill=FILL_MINT, border=True, italic=True, color=C_MUTED, align="right")
    put(ws, f"B{t + 1}", "Still to pay", bold=True, border=True)
    for mi in range(12):
        L = get_column_letter(6 + mi)
        calc_cell(ws, f"{L}{t}", f'=SUMIF({L}{first}:{L}{last},"Paid",$D${first}:$D${last})', fmt=MONEY0,
                  bold=True, fill=FILL_MINT, align="center")
        calc_cell(ws, f"{L}{t + 1}", f"=$D${t}-{L}{t}", fmt=MONEY0, align="center")
    ws.conditional_formatting.add(f"F{t + 1}:Q{t + 1}", CellIsRule(operator="greaterThan", formula=["0"],
                                                                   font=Font(name=FONT, color=C_RED, bold=True)))
    ws.freeze_panes = "C6"
    print_setup(ws)
    return ws


def build_savings(wb, demo):
    ws = wb.create_sheet(S_SAVE)
    title_block(ws, "Savings Goals", "Set a target and a date — see exactly how much to save each month.", 9)
    set_widths(ws, {"B": 28, "C": 13, "D": 13, "E": 14, "F": 9, "G": 16, "H": 13, "I": 11, "J": 14})
    header_row(ws, 5, 2, ["Goal", "Target", "Saved so far", "Target date", "%", "Progress",
                          "Remaining", "Months left", "Save / month"], height=26)
    first, last = 6, 20
    for i, r in enumerate(range(first, last + 1)):
        g = DEMO_GOALS[i] if demo and i < len(DEMO_GOALS) else None
        input_cell(ws, f"B{r}", g[0] if g else None)
        input_cell(ws, f"C{r}", g[1] if g else None, fmt=MONEY)
        input_cell(ws, f"D{r}", g[2] if g else None, fmt=MONEY)
        input_cell(ws, f"E{r}", g[3] if g else None, fmt=DATE_FMT, align="center")
        calc_cell(ws, f"F{r}", f'=IF(B{r}="","",IFERROR(MIN(1,D{r}/C{r}),0))', fmt=PCT, align="center")
        calc_cell(ws, f"G{r}", f'=IF(B{r}="","",{progress_bar(f"F{r}")})', color=C_PRIMARY)
        calc_cell(ws, f"H{r}", f'=IF(B{r}="","",MAX(0,N(C{r})-N(D{r})))', fmt=MONEY)
        calc_cell(ws, f"I{r}", f'=IF(OR(B{r}="",E{r}=""),"",MAX(0,(YEAR(E{r})-YEAR(TODAY()))*12+MONTH(E{r})-MONTH(TODAY())))',
                  fmt="0", align="center")
        calc_cell(ws, f"J{r}", f'=IF(I{r}="","",IF(H{r}=0,0,H{r}/MAX(1,I{r})))', fmt=MONEY, bold=True)
    t = last + 1
    put(ws, f"B{t}", "Total", bold=True, fill=FILL_MINT, border=True)
    for col in "CDHJ":
        calc_cell(ws, f"{col}{t}", f"=SUM({col}{first}:{col}{last})", fmt=MONEY, bold=True, fill=FILL_MINT)
    calc_cell(ws, f"F{t}", f"=IFERROR(D{t}/C{t},0)", fmt=PCT, bold=True, fill=FILL_MINT, align="center")
    calc_cell(ws, f"G{t}", f'={progress_bar(f"F{t}")}', bold=True, fill=FILL_MINT, color=C_PRIMARY)
    for col in "EI":
        put(ws, f"{col}{t}", "", fill=FILL_MINT, border=True)
    ws.conditional_formatting.add(f"F{first}:F{last}", CellIsRule(operator="greaterThanOrEqual", formula=["1"],
                                                                  fill=FILL_GREEN_LIGHT, font=Font(name=FONT, color="2E7D4F", bold=True)))
    put(ws, f"B{t + 2}", "Tip: 'Save / month' spreads what is left over the months until your target date.",
        size=8, italic=True, color=C_MUTED)

    bar = BarChart()
    bar.type = "bar"
    bar.grouping = "clustered"
    bar.add_data(Reference(ws, min_col=3, max_col=4, min_row=5, max_row=5 + max(len(DEMO_GOALS), 5)), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=2, min_row=6, max_row=5 + max(len(DEMO_GOALS), 5)))
    style_chart(bar, "Target vs. saved", width=18, height=7.5)
    color_series(bar, [C_SAGE, C_PRIMARY])
    ws.add_chart(bar, f"B{t + 4}")
    ws.freeze_panes = "C6"
    print_setup(ws)
    return ws


def build_networth(wb, demo):
    ws = wb.create_sheet(S_NW)
    title_block(ws, "Net Worth Tracker", "Update balances once a month (e.g. on the 1st) and watch your net worth grow.", 13)
    set_widths(ws, {"B": 26, **{get_column_letter(c): 11 for c in range(3, 15)}})
    header_row(ws, 5, 2, ["Account"] + MON3)

    def block(label_row, rows, demo_rows, label):
        put(ws, f"B{label_row}", label, bold=True, color=C_WHITE, fill=FILL_SAGE)
        for col in range(3, 15):
            ws.cell(row=label_row, column=col).fill = FILL_SAGE
        first = label_row + 1
        for i in range(rows):
            r = first + i
            d = demo_rows[i] if demo and i < len(demo_rows) else None
            input_cell(ws, f"B{r}", d[0] if d else None)
            for mi in range(12):
                L = get_column_letter(3 + mi)
                input_cell(ws, f"{L}{r}", round(d[1] + d[2] * mi, 2) if d and mi < 12 else None, fmt=MONEY0)
        last = first + rows - 1
        tot = last + 1
        put(ws, f"B{tot}", f"Total {label.lower()}", bold=True, fill=FILL_MINT, border=True)
        for mi in range(12):
            L = get_column_letter(3 + mi)
            calc_cell(ws, f"{L}{tot}", f"=SUM({L}{first}:{L}{last})", fmt=MONEY0, bold=True, fill=FILL_MINT)
        return tot

    a_tot = block(6, 10, DEMO_ASSETS, "ASSETS")
    l_tot = block(a_tot + 2, 10, DEMO_LIABS, "LIABILITIES")
    nw = l_tot + 2
    put(ws, f"B{nw}", "NET WORTH", bold=True, fill=FILL_GOLD_LIGHT, border=True, size=11)
    put(ws, f"B{nw + 1}", "Change vs. last month", bold=True, border=True)
    for mi in range(12):
        L = get_column_letter(3 + mi)
        calc_cell(ws, f"{L}{nw}", f"={L}{a_tot}-{L}{l_tot}", fmt=MONEY0, bold=True, fill=FILL_GOLD_LIGHT)
        if mi == 0:
            put(ws, f"{L}{nw + 1}", "—", border=True, align="center", color=C_MUTED)
        else:
            P = get_column_letter(2 + mi)
            calc_cell(ws, f"{L}{nw + 1}", f'=IF(COUNT({L}7:{L}{l_tot})=0,"",{L}{nw}-{P}{nw})', fmt=MONEY0)
    ws.conditional_formatting.add(f"C{nw}:N{nw + 1}", CellIsRule(operator="lessThan", formula=["0"],
                                                                 font=Font(name=FONT, color=C_RED, bold=True)))
    line = LineChart()
    line.add_data(Reference(ws, min_col=2, max_col=14, min_row=nw), from_rows=True, titles_from_data=True)
    line.set_categories(Reference(ws, min_col=3, max_col=14, min_row=5))
    style_chart(line, "Net worth over the year", width=22, height=7.5)
    line.series[0].graphicalProperties.line.solidFill = C_PRIMARY
    line.series[0].graphicalProperties.line.width = 32000
    line.series[0].smooth = False
    ws.add_chart(line, f"B{nw + 3}")
    ws.freeze_panes = "C6"
    print_setup(ws)
    return ws


def build_debt(wb, demo):
    ws = wb.create_sheet(S_DEBT)
    title_block(ws, "Debt Payoff Planner", "Snowball = smallest balance first. Avalanche = highest interest first.", 15)
    set_widths(ws, {"B": 24, "C": 14, "D": 10, "E": 14, "F": 10, "G": 12, "H": 13, "I": 14, "J": 14,
                    "K": 3, "L": 8, "M": 20, "N": 12, "O": 9, "P": 12})

    put(ws, "B6", "Strategy", bold=True, border=True, fill=FILL_MINT)
    input_cell(ws, "C6", "Snowball", align="center")
    add_list_validation(ws, "C6", f"={LISTS}!$F$2:$F$3", "Snowball or Avalanche")
    put(ws, "B7", "Extra payment / month", bold=True, border=True, fill=FILL_MINT)
    input_cell(ws, "C7", 200 if demo else 0, fmt=MONEY)
    put(ws, "B8", "First payment month", bold=True, border=True, fill=FILL_MINT)
    if demo:
        input_cell(ws, "C8", dt.date(YEAR, 1, 1), fmt=MONTH_FMT, align="center")
    else:
        input_cell(ws, "C8", "=DATE(YEAR(TODAY()),MONTH(TODAY())+1,1)", fmt=MONTH_FMT, align="center")
    put(ws, "B9", "Total paid toward debt / month", bold=True, border=True, fill=FILL_MINT)
    calc_cell(ws, "C9", f"=SUM(E{DEBT_FIRST}:E{DEBT_LAST})+N(C7)", fmt=MONEY, bold=True)
    put(ws, "D6", "← pick one", size=8, italic=True, color=C_MUTED)
    put(ws, "D7", "← anything above the minimums", size=8, italic=True, color=C_MUTED)

    # KPI cards
    kpi_card(ws, 6, 5, "DEBT-FREE DATE",
             f'=IF(COUNT(C{DEBT_FIRST}:C{DEBT_LAST})=0,"—",IF(COUNTIF(G{DEBT_FIRST}:G{DEBT_LAST},"30+ yrs")>0,"Pay more",'
             f'DATE(YEAR(C8),MONTH(C8)+MAX(G{DEBT_FIRST}:G{DEBT_LAST})-1,1)))', fmt=MONTH_FMT, width_cols=2)
    ws["F7"].alignment = Alignment(horizontal="center", vertical="center")
    kpi_card(ws, 8, 5, "MONTHS TO GO", f'=IF(COUNT(G{DEBT_FIRST}:G{DEBT_LAST})=0,"—",MAX(G{DEBT_FIRST}:G{DEBT_LAST}))',
             fmt="0", width_cols=2)
    kpi_card(ws, 10, 5, "TOTAL INTEREST", f"=SUM(I{DEBT_FIRST}:I{DEBT_LAST})", width_cols=1)
    ws.column_dimensions["J"].width = 14

    header_row(ws, DEBT_FIRST - 1, 2, ["Debt", "Balance", "APR (%)", "Min. payment", "Priority",
                                       "Paid off in (months)", "Payoff date", "Interest paid", "Total paid"], height=30)
    for i, r in enumerate(range(DEBT_FIRST, DEBT_LAST + 1)):
        d = DEMO_DEBTS[i] if demo and i < len(DEMO_DEBTS) else None
        input_cell(ws, f"B{r}", d[0] if d else None)
        input_cell(ws, f"C{r}", d[1] if d else None, fmt=MONEY)
        input_cell(ws, f"D{r}", d[2] if d else None, fmt="0.00", align="center")
        input_cell(ws, f"E{r}", d[3] if d else None, fmt=MONEY)
        calc_cell(ws, f"F{r}",
                  f'=IF(OR(B{r}="",N(C{r})<=0),"",IFERROR(IF($C$6="Avalanche",'
                  f'RANK(D{r},$D${DEBT_FIRST}:$D${DEBT_LAST},0)+COUNTIF($D${DEBT_FIRST - 1}:D{r - 1},D{r}),'
                  f'RANK(C{r},$C${DEBT_FIRST}:$C${DEBT_LAST},1)+COUNTIF($C${DEBT_FIRST - 1}:C{r - 1},C{r})),""))',
                  fmt="0", align="center", bold=True)
        bal_cols = f"{SCHED}!$F${SCHED_FIRST}:$K${SCHED_LAST}"
        pay_cols = f"{SCHED}!$L${SCHED_FIRST}:$Q${SCHED_LAST}"
        cnt = f'COUNTIF(INDEX({bal_cols},0,F{r}),">0.004")'
        calc_cell(ws, f"G{r}", f'=IF(F{r}="","",IF({cnt}>={SCHED_MONTHS},"30+ yrs",{cnt}+1))', fmt="0", align="center")
        calc_cell(ws, f"H{r}", f'=IF(ISNUMBER(G{r}),DATE(YEAR($C$8),MONTH($C$8)+G{r}-1,1),"")', fmt=MONTH_FMT, align="center")
        calc_cell(ws, f"J{r}", f'=IF(F{r}="","",SUM(INDEX({pay_cols},0,F{r})))', fmt=MONEY)
        calc_cell(ws, f"I{r}", f'=IF(F{r}="","",MAX(0,J{r}-C{r}))', fmt=MONEY)
    t = DEBT_LAST + 1
    put(ws, f"B{t}", "Total", bold=True, fill=FILL_MINT, border=True)
    for col, fmt in (("C", MONEY), ("E", MONEY), ("I", MONEY), ("J", MONEY)):
        calc_cell(ws, f"{col}{t}", f"=SUM({col}{DEBT_FIRST}:{col}{DEBT_LAST})", fmt=fmt, bold=True, fill=FILL_MINT)
    for col in "DFGH":
        put(ws, f"{col}{t}", "", fill=FILL_MINT, border=True)
    put(ws, f"B{t + 1}", "Tip: if a debt shows '30+ yrs', its minimum payment doesn't cover the interest — add an extra payment.",
        size=8, italic=True, color=C_MUTED)
    put(ws, f"B{t + 2}", "APR: type 19.99 for 19.99%. Priority 1 gets every spare dollar first; "
                         "when a debt is paid off, its payment rolls to the next one.", size=8, italic=True, color=C_MUTED)

    # Sorted helper table (payoff order) feeding the schedule
    put(ws, f"L{DEBT_FIRST - 2}", "PAYOFF ORDER", size=11, bold=True, color=C_PRIMARY)
    header_row(ws, DEBT_FIRST - 1, 12, ["#", "Debt", "Balance", "APR", "Min."], height=30)
    for k in range(1, 7):
        r = DEBT_FIRST + k - 1
        put(ws, f"L{r}", k, bold=True, border=True, align="center", fill=FILL_MINT)
        mt = f"MATCH($L{r},$F${DEBT_FIRST}:$F${DEBT_LAST},0)"
        calc_cell(ws, f"M{r}", f'=IFERROR(INDEX($B${DEBT_FIRST}:$B${DEBT_LAST},{mt}),"")')
        calc_cell(ws, f"N{r}", f'=IFERROR(N(INDEX($C${DEBT_FIRST}:$C${DEBT_LAST},{mt})),0)', fmt=MONEY0)
        calc_cell(ws, f"O{r}", f'=IFERROR(N(INDEX($D${DEBT_FIRST}:$D${DEBT_LAST},{mt})),0)', fmt="0.00", align="center")
        calc_cell(ws, f"P{r}", f'=IFERROR(N(INDEX($E${DEBT_FIRST}:$E${DEBT_LAST},{mt})),0)', fmt=MONEY0)

    ws.freeze_panes = "A5"
    print_setup(ws)
    return ws


def build_schedule(wb, debt_ws):
    ws = wb.create_sheet(S_SCHED)
    title_block(ws, "Debt Schedule", "Month-by-month plan (auto). Grey columns on the right are the math behind it.", 16)
    set_widths(ws, {"B": 8, "C": 11, "D": 12, "E": 13, **{get_column_letter(c): 12 for c in range(6, 18)}})
    for c in range(18, 36):
        ws.column_dimensions[get_column_letter(c)].width = 10
    D = DEBT
    put(ws, "F5", "BALANCE AT MONTH END", bold=True, color=C_PRIMARY)
    put(ws, "L5", "PAYMENT THIS MONTH", bold=True, color=C_PRIMARY)
    put(ws, "R5", "helper: owed after interest", size=8, italic=True, color=C_MUTED)
    put(ws, "X5", "helper: minimum part", size=8, italic=True, color=C_MUTED)
    put(ws, "AD5", "helper: extra part", size=8, italic=True, color=C_MUTED)
    labels = ["Month #", "Month", "Total paid", "Total balance"]
    header_row(ws, SCHED_START_ROW - 1, 2, labels, height=30)
    for block_start, _ in ((6, "bal"), (12, "pay"), (18, "owed"), (24, "min"), (30, "extra")):
        for k in range(6):
            col = get_column_letter(block_start + k)
            src = f"{D}!$M${DEBT_FIRST + k}"
            c = ws[f"{col}{SCHED_START_ROW - 1}"]
            c.value = f'=IF({src}="","Debt {k + 1}",{src})'
            grey = block_start >= 18
            c.font = font(9, True, C_WHITE if not grey else C_TEXT)
            c.fill = FILL_PRIMARY if not grey else FILL_GREY_LIGHT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER

    r0 = SCHED_START_ROW
    put(ws, f"B{r0}", 0, border=True, align="center", color=C_MUTED)
    put(ws, f"C{r0}", "Start", border=True, align="center", color=C_MUTED)
    put(ws, f"D{r0}", "", border=True)
    calc_cell(ws, f"E{r0}", f"=SUM(F{r0}:K{r0})", fmt=MONEY0, bold=True)
    for k in range(6):
        col = get_column_letter(6 + k)
        calc_cell(ws, f"{col}{r0}", f"={D}!$N${DEBT_FIRST + k}", fmt=MONEY0)

    budget = f"{D}!$C$9"
    grey_font = font(9, color=C_MUTED)
    for m in range(1, SCHED_MONTHS + 1):
        r = r0 + m
        ws[f"B{r}"] = m
        ws[f"C{r}"] = f"=DATE(YEAR({D}!$C$8),MONTH({D}!$C$8)+B{r}-1,1)"
        ws[f"D{r}"] = f"=SUM(L{r}:Q{r})"
        ws[f"E{r}"] = f"=SUM(F{r}:K{r})"
        for k in range(6):
            bal, pay, owed, mn, ex = (get_column_letter(6 + k), get_column_letter(12 + k),
                                      get_column_letter(18 + k), get_column_letter(24 + k),
                                      get_column_letter(30 + k))
            apr = f"{D}!$O${DEBT_FIRST + k}"
            minp = f"{D}!$P${DEBT_FIRST + k}"
            ws[f"{owed}{r}"] = f"=IF({bal}{r - 1}<=0.004,0,ROUND({bal}{r - 1}*(1+{apr}/1200),2))"
            ws[f"{mn}{r}"] = f"=MIN({minp},{owed}{r})"
            prior_extra = "" if k == 0 else f"-SUM($AD{r}:{get_column_letter(30 + k - 1)}{r})"
            ws[f"{ex}{r}"] = f"=MAX(0,MIN({owed}{r}-{mn}{r},{budget}-SUM($X{r}:$AC{r}){prior_extra}))"
            ws[f"{pay}{r}"] = f"={mn}{r}+{ex}{r}"
            ws[f"{bal}{r}"] = f"=MAX(0,ROUND({owed}{r}-{pay}{r},2))"
        for c in range(2, 36):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            if c == 2:
                cell.alignment = Alignment(horizontal="center")
                cell.font = font(9, color=C_MUTED)
            elif c == 3:
                cell.number_format = MONTH_FMT
                cell.alignment = Alignment(horizontal="center")
                cell.font = font(9)
            elif c >= 18:
                cell.number_format = MONEY0
                cell.font = grey_font
            else:
                cell.number_format = MONEY0
                cell.font = font(9, bold=(c == 5))
    ws.conditional_formatting.add(f"F{SCHED_FIRST}:K{SCHED_LAST}",
                                  FormulaRule(formula=[f"AND(F{SCHED_FIRST}=0,F{SCHED_FIRST - 1}>0)"],
                                              fill=FILL_GREEN_LIGHT, font=Font(name=FONT, color="2E7D4F", bold=True)))
    ws.freeze_panes = f"D{SCHED_FIRST}"

    # Balance-over-time chart placed on the Debt Payoff sheet
    line = LineChart()
    line.add_data(Reference(ws, min_col=5, min_row=SCHED_START_ROW - 1, max_row=SCHED_START_ROW + 72), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=2, min_row=SCHED_START_ROW, max_row=SCHED_START_ROW + 72))
    style_chart(line, "Total debt balance — first 6 years", width=24, height=7.5)
    line.series[0].graphicalProperties.line.solidFill = C_RED
    line.series[0].graphicalProperties.line.width = 32000
    line.legend = None
    line.x_axis.title = "Month #"
    line.x_axis.tickLblSkip = 6
    debt_ws.add_chart(line, f"B{DEBT_LAST + 5}")
    print_setup(ws)
    return ws


def build_holiday(wb, demo):
    ws = wb.create_sheet(S_HOL)
    title_block(ws, "Holiday Budget & Gift Tracker", "Who you're buying for, what, how much — and what's still left to do.", 15)
    set_widths(ws, {"B": 22, "C": 11, "D": 11, "E": 24, "F": 16, "G": 12, "H": 11, "I": 12, "J": 20,
                    "K": 3, "L": 20, "M": 12, "N": 12, "O": 12})

    put(ws, "B6", "Holiday date", bold=True, border=True, fill=FILL_MINT)
    if demo:
        input_cell(ws, "C6", dt.date(2026, 12, 25), fmt=DATE_FMT, align="center")
    else:
        input_cell(ws, "C6", "=DATE(YEAR(TODAY())+IF(TODAY()>DATE(YEAR(TODAY()),12,25),1,0),12,25)",
                   fmt=DATE_FMT, align="center")
    ws.merge_cells("C6:D6")
    put(ws, "B7", "Total holiday budget", bold=True, border=True, fill=FILL_MINT)
    input_cell(ws, "C7", 1500 if demo else None, fmt=MONEY)
    ws.merge_cells("C7:D7")
    put(ws, "B8", "Days to go", bold=True, border=True, fill=FILL_MINT)
    calc_cell(ws, "C8", "=MAX(0,C6-TODAY())", fmt="0", bold=True, align="center")
    ws.merge_cells("C8:D8")

    gf, gl = 12, 51
    of, ol = 20, 28
    gifts_budget = f"SUM(D{gf}:D{gl})"
    gifts_spent = f"SUM(G{gf}:G{gl})"
    other_spent = f"SUM(N{of}:N{ol})"
    kpi_card(ws, 5, 5, "GIFTS BOUGHT", f'=IFERROR((COUNTIF(H{gf}:H{gl},"Bought")+COUNTIF(H{gf}:H{gl},"Wrapped")'
                                         f'+COUNTIF(H{gf}:H{gl},"Given"))/COUNTA(B{gf}:B{gl}),0)', fmt=PCT)
    ws["E7"] = f'=COUNTA(B{gf}:B{gl})&" people on the list"'
    kpi_card(ws, 7, 5, "SPENT SO FAR", f"={gifts_spent}+{other_spent}")
    ws["G7"] = f'="gifts "&TEXT({gifts_spent},"#,##0")&" + other "&TEXT({other_spent},"#,##0")'
    kpi_card(ws, 9, 5, "LEFT IN BUDGET", f"=N(C7)-({gifts_spent}+{other_spent})")
    ws["I7"] = f'={progress_bar(f"IFERROR(({gifts_spent}+{other_spent})/C7,0)")}&" used"'
    for ref in ("E7", "G7", "I7"):
        ws[ref].font = font(8, italic=True, color=C_MUTED)
        ws[ref].alignment = Alignment(horizontal="center")
    ws.conditional_formatting.add("I6", CellIsRule(operator="lessThan", formula=["0"],
                                                   font=Font(name=FONT, color=C_RED, bold=True, size=16)))

    put(ws, f"B{gf - 2}", "GIFT LIST", size=12, bold=True, color=C_PRIMARY)
    header_row(ws, gf - 1, 2, ["Recipient", "Group", "Budget", "Gift idea", "Store / link", "Actual cost",
                               "Status", "Difference", "Notes"])
    for i, r in enumerate(range(gf, gl + 1)):
        g = DEMO_GIFTS[i] if demo and i < len(DEMO_GIFTS) else None
        input_cell(ws, f"B{r}", g[0] if g else None)
        input_cell(ws, f"C{r}", g[1] if g else None, align="center")
        input_cell(ws, f"D{r}", g[2] if g else None, fmt=MONEY)
        input_cell(ws, f"E{r}", g[3] if g else None)
        input_cell(ws, f"F{r}", g[4] if g else None)
        input_cell(ws, f"G{r}", (g[5] or None) if g else None, fmt=MONEY)
        input_cell(ws, f"H{r}", g[6] if g else None, align="center")
        calc_cell(ws, f"I{r}", f'=IF(OR(B{r}="",G{r}=""),"",D{r}-G{r})', fmt=MONEY)
        input_cell(ws, f"J{r}")
    add_list_validation(ws, f"C{gf}:C{gl}", f"={LISTS}!$E$2:$E$6")
    add_list_validation(ws, f"H{gf}:H{gl}", f"={LISTS}!$D$2:$D$6", "Idea → Ordered → Bought → Wrapped → Given")
    status_colors = {"Idea": FILL_GREY_LIGHT, "Ordered": FILL_BLUE_LIGHT, "Bought": FILL_GOLD_LIGHT,
                     "Wrapped": PatternFill("solid", fgColor="E8DDF3"), "Given": FILL_GREEN_LIGHT}
    for s, fill in status_colors.items():
        ws.conditional_formatting.add(f"H{gf}:H{gl}", CellIsRule(operator="equal", formula=[f'"{s}"'], fill=fill,
                                                                 font=Font(name=FONT, bold=True)))
    ws.conditional_formatting.add(f"I{gf}:I{gl}", CellIsRule(operator="lessThan", formula=["0"],
                                                             font=Font(name=FONT, color=C_RED, bold=True)))
    t = gl + 1
    put(ws, f"B{t}", "Total gifts", bold=True, fill=FILL_MINT, border=True)
    for col in "CEFHJ":
        put(ws, f"{col}{t}", "", fill=FILL_MINT, border=True)
    for col in "DGI":
        calc_cell(ws, f"{col}{t}", f"=SUM({col}{gf}:{col}{gl})", fmt=MONEY, bold=True, fill=FILL_MINT)

    # Status summary
    put(ws, "L10", "STATUS", size=12, bold=True, color=C_PRIMARY)
    header_row(ws, 11, 12, ["Status", "Gifts"])
    for i, s in enumerate(status_colors):
        r = 12 + i
        put(ws, f"L{r}", s, bold=True, border=True, fill=status_colors[s])
        calc_cell(ws, f"M{r}", f'=COUNTIF($H${gf}:$H${gl},L{r})', fmt="0", align="center")

    # Other holiday costs
    put(ws, f"L{of - 2}", "OTHER HOLIDAY COSTS", size=12, bold=True, color=C_PRIMARY)
    header_row(ws, of - 1, 12, ["Category", "Budget", "Actual", "Difference"])
    for i, r in enumerate(range(of, ol + 1)):
        name = HOLIDAY_OTHER[i]
        input_cell(ws, f"L{r}", name)
        b, a = DEMO_HOLIDAY_OTHER[name] if demo else (None, None)
        input_cell(ws, f"M{r}", b, fmt=MONEY)
        input_cell(ws, f"N{r}", a if a else None, fmt=MONEY)
        calc_cell(ws, f"O{r}", f'=IF(L{r}="","",N(M{r})-N(N{r}))', fmt=MONEY)
    ot = ol + 1
    put(ws, f"L{ot}", "Total other", bold=True, fill=FILL_MINT, border=True)
    for col in "MNO":
        calc_cell(ws, f"{col}{ot}", f"=SUM({col}{of}:{col}{ol})", fmt=MONEY, bold=True, fill=FILL_MINT)
    ws.conditional_formatting.add(f"O{of}:O{ot}", CellIsRule(operator="lessThan", formula=["0"],
                                                             font=Font(name=FONT, color=C_RED, bold=True)))

    # Budget check
    bc = ot + 2
    put(ws, f"L{bc}", "BUDGET CHECK", size=12, bold=True, color=C_PRIMARY)
    rows = [("Gift budgets", f"=D{t}"), ("Other budgets", f"=M{ot}"), ("Planned total", f"=M{bc + 1}+M{bc + 2}"),
            ("Your total budget", "=N(C7)"), ("Unplanned", f"=M{bc + 4}-M{bc + 3}")]
    for i, (lab, f) in enumerate(rows):
        r = bc + 1 + i
        put(ws, f"L{r}", lab, bold=True, border=True, fill=FILL_MINT if i in (2, 4) else None)
        calc_cell(ws, f"M{r}", f, fmt=MONEY, bold=i in (2, 4), fill=FILL_MINT if i in (2, 4) else None)
    ws.conditional_formatting.add(f"M{bc + 5}", CellIsRule(operator="lessThan", formula=["0"],
                                                           font=Font(name=FONT, color=C_RED, bold=True)))

    # Spend by group
    sg = bc + 8
    put(ws, f"L{sg}", "GIFT SPEND BY GROUP", size=12, bold=True, color=C_PRIMARY)
    header_row(ws, sg + 1, 12, ["Group", "Budget", "Spent"])
    for i, g in enumerate(["Family", "Friends", "Kids", "Work", "Other"]):
        r = sg + 2 + i
        put(ws, f"L{r}", g, bold=True, border=True, fill=FILL_MINT)
        calc_cell(ws, f"M{r}", f"=SUMIF($C${gf}:$C${gl},L{r},$D${gf}:$D${gl})", fmt=MONEY)
        calc_cell(ws, f"N{r}", f"=SUMIF($C${gf}:$C${gl},L{r},$G${gf}:$G${gl})", fmt=MONEY)
    bar = BarChart()
    bar.type = "bar"
    bar.add_data(Reference(ws, min_col=13, max_col=14, min_row=sg + 1, max_row=sg + 6), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=12, min_row=sg + 2, max_row=sg + 6))
    style_chart(bar, "Gift budget vs. spent by group", width=14, height=7)
    color_series(bar, [C_SAGE, C_RED])
    ws.add_chart(bar, f"L{sg + 8}")
    ws.freeze_panes = "A5"
    print_setup(ws)
    return ws


# --------------------------------------------------------------------------
EDITIONS = {
    "bundle": {
        "file": "2027-Money-Planner",
        "tabs": ["budget", "bills", "savings", "networth", "debt", "holiday"],
    },
    "holiday": {"file": "Holiday-Budget-Gift-Tracker", "tabs": ["holiday"]},
    "debt": {"file": "Debt-Payoff-Planner", "tabs": ["debt"]},
}


# Print areas used only when rendering preview screenshots (one page per sheet)
RENDER_AREAS = {
    S_START: "A1:G50", S_CAT: "A1:I52", S_TX: "A1:H42", S_MONTH: "A1:M65", S_ANNUAL: "A1:R75",
    S_BILLS: "A1:Q33", S_SAVE: "A1:J40", S_NW: "A1:N50", S_DEBT: "A1:P38", S_SCHED: "A1:Q48",
    S_HOL: "A1:O70",
}


def build(edition, demo, outdir, render=False):
    tabs = EDITIONS[edition]["tabs"]
    PLANNED_SHEETS.clear()
    if "budget" in tabs:
        PLANNED_SHEETS.update({S_CAT, S_TX, S_MONTH, S_ANNUAL})
    if "bills" in tabs:
        PLANNED_SHEETS.add(S_BILLS)
    if "savings" in tabs:
        PLANNED_SHEETS.add(S_SAVE)
    if "networth" in tabs:
        PLANNED_SHEETS.add(S_NW)
    if "debt" in tabs:
        PLANNED_SHEETS.update({S_DEBT, S_SCHED})
    if "holiday" in tabs:
        PLANNED_SHEETS.add(S_HOL)

    wb = Workbook()
    build_start(wb, edition, tabs)
    if "budget" in tabs:
        build_categories(wb, demo)
        build_transactions(wb, demo)
        build_monthly(wb)
        build_annual(wb)
    if "bills" in tabs:
        build_bills(wb, demo)
    if "savings" in tabs:
        build_savings(wb, demo)
    if "networth" in tabs:
        build_networth(wb, demo)
    if "debt" in tabs:
        debt_ws = build_debt(wb, demo)
        build_schedule(wb, debt_ws)
    if "holiday" in tabs:
        build_holiday(wb, demo)
    build_lists(wb)

    # Tab colors
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = {S_START: C_GOLD, S_HOL: C_RED}.get(ws.title, C_PRIMARY)
        if ws.title in (S_TX, S_CAT):
            ws.sheet_properties.tabColor = C_SAGE
    wb.calculation.fullCalcOnLoad = True
    if render:
        for ws in wb.worksheets:
            if ws.title in RENDER_AREAS:
                ws.print_area = RENDER_AREAS[ws.title]
                ws.page_setup.fitToHeight = 1
                ws.page_setup.orientation = "landscape"
    name = EDITIONS[edition]["file"] + ("-DEMO" if demo else "") + ("-RENDER" if render else "") + ".xlsx"
    path = os.path.join(outdir, name)
    wb.save(path)
    return path


def main():
    ap = argparse.ArgumentParser()
    here = os.path.dirname(os.path.abspath(__file__))
    ap.add_argument("--outdir", default=os.path.join(here, "..", "dist"))
    ap.add_argument("--edition", choices=list(EDITIONS) + ["all"], default="all")
    args = ap.parse_args()
    os.makedirs(args.outdir, exist_ok=True)
    editions = list(EDITIONS) if args.edition == "all" else [args.edition]
    for e in editions:
        for demo in (False, True):
            print(build(e, demo, args.outdir))


if __name__ == "__main__":
    main()
